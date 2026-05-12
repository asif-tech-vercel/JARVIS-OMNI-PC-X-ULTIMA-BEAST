"""JARVIS Document Writer Skill."""

try:
    from docx import Document
    from docx.shared import Inches, Pt
    HAS_DOCX = True
except: HAS_DOCX = False

try:
    from openpyxl import Workbook
    HAS_XLSX = True
except: HAS_XLSX = False

from jarvis_core.logger import get_logger
logger = get_logger()

def create_word_doc(path: str, content: str, title: str = None) -> dict:
    """Create Word document."""
    if not HAS_DOCX:
        return {"success": False, "error": "python-docx not available"}
    try:
        doc = Document()
        if title:
            doc.add_heading(title, 0)
        doc.add_paragraph(content)
        doc.save(path)
        logger.action(f"Created Word doc: {path}", "SUCCESS", "low")
        return {"success": True, "path": path}
    except Exception as e:
        return {"success": False, "error": str(e)}

def create_excel(path: str, data: list, headers: list = None) -> dict:
    """Create Excel file."""
    if not HAS_XLSX:
        return {"success": False, "error": "openpyxl not available"}
    try:
        wb = Workbook()
        ws = wb.active
        if headers:
            ws.append(headers)
        for row in data:
            ws.append(row)
        wb.save(path)
        logger.action(f"Created Excel: {path}", "SUCCESS", "low")
        return {"success": True, "path": path}
    except Exception as e:
        return {"success": False, "error": str(e)}